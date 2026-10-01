"""Build the project presentation. Run AFTER training and after saving screenshots.

    python build_ppt.py --name "Your Name" --reg "RA24xxxxxxxxx" --course "Artificial Neural Networks"

Reads artifacts/metrics.json for the numbers and screenshots/*.png for the demo slides.
"""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
BG, CARD, LINE = "0B1220", "111A2E", "1E2A44"
TEXT, MUTED, TEAL, CORAL, AMBER = "E6EAF2", "8B97AD", "2DD4BF", "FB7185", "FBBF24"
FONT = "Calibri"

SHOTS = [
    ("01_analyze.png", "Live analysis", "Gauge score and word-level highlighting update as the user types."),
    ("02_negative.png", "Negative review", "The same network flags the words driving a negative verdict."),
    ("03_insights.png", "Model insights", "Training curves, confusion matrix and architecture in one place."),
    ("04_how_it_works.png", "How it works", "The six-stage pipeline explained for non-technical viewers."),
]


def rgb(h):
    return RGBColor.from_string(h)


def box(slide, x, y, w, h, fill=CARD, line=None, round_=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        shp.adjustments[0] = 0.05
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, content, size=16, bold=False, color=TEXT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = rgb(color)
    return tb


def new_slide(prs, title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    text(s, 0.7, 0.5, 11.9, 0.7, title, size=32, bold=True)
    if subtitle:
        text(s, 0.7, 1.15, 11.9, 0.4, subtitle, size=15, color=MUTED)
    return s


def bullets(slide, x, y, w, items, size=18, gap=14):
    for i, (head, body) in enumerate(items):
        text(slide, x, y + i * (0.55 + gap / 72 * 4), w, 0.4, head, size=size, bold=True, color=TEAL)
        text(slide, x, y + i * (0.55 + gap / 72 * 4) + 0.33, w, 0.6, body, size=size - 3, color=MUTED)


def picture_or_placeholder(slide, path, x, y, w, h):
    box(slide, x, y, w, h, fill=CARD, line=LINE)
    if path.exists():
        from PIL import Image

        iw, ih = Image.open(path).size
        scale = min((w - 0.2) / iw, (h - 0.2) / ih)
        pw, ph = iw * scale, ih * scale
        slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                 Inches(pw), Inches(ph))
    else:
        text(slide, x, y, w, h, f"Save screenshot as screenshots/{path.name} and rebuild",
             size=16, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--name", default="Your Name")
    ap.add_argument("--reg", default="Register No.")
    ap.add_argument("--course", default="Artificial Neural Networks")
    ap.add_argument("--github", default="github.com/your-username/sentiscope")
    ap.add_argument("--demo", default="your-app.streamlit.app")
    ap.add_argument("--out", default="SentiScope_ANN_Activity2.pptx")
    a = ap.parse_args()

    mp = ROOT / "artifacts" / "metrics.json"
    m = json.loads(mp.read_text()) if mp.exists() else None
    acc = f"{m['test_acc'] * 100:.1f}%" if m else "-"
    params = f"{m['params']:,}" if m else "-"
    epochs = str(m["epochs"]) if m else "-"

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # 1 title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    text(s, 0.9, 2.2, 11, 1.2, "SentiScope", size=64, bold=True)
    text(s, 0.9, 3.45, 11, 0.8, "Real-time sentiment analysis with a neural network that explains every word",
         size=22, color=MUTED)
    text(s, 0.9, 5.3, 11, 1.2, [a.name + "  |  " + a.reg, a.course + "  |  Activity 2: Interactive Neural Network Demo"],
         size=16, color=TEXT, spacing=4)

    # 2 problem
    s = new_slide(prs, "Problem and objective", "Why this demo exists")
    bullets(s, 0.9, 1.9, 5.6, [
        ("The problem", "Neural networks feel like black boxes. A score alone does not build trust."),
        ("The objective", "Build an interactive web app that predicts sentiment live and shows why."),
        ("The deliverable", "A deployed Streamlit app, source on GitHub, with a recorded walkthrough."),
    ])
    box(s, 7.2, 1.9, 5.4, 4.6)
    text(s, 7.6, 2.2, 4.6, 0.4, "WHAT THE USER CAN DO", size=12, bold=True, color=MUTED)
    text(s, 7.6, 2.7, 4.6, 3.6, ["Type text and watch the score move", "See which words pushed it positive or negative",
                                 "Inspect training curves and the confusion matrix", "Learn the pipeline in plain language"],
         size=17, spacing=14)

    # 3 dataset
    s = new_slide(prs, "Dataset and preprocessing", "IMDB Large Movie Review Dataset")
    for i, (v, l) in enumerate([("50,000", "Labelled reviews"), ("25,000 / 25,000", "Train / test split"),
                                ("10,000", "Vocabulary size"), ("200", "Tokens per review")]):
        box(s, 0.7 + i * 3.05, 1.9, 2.85, 1.5)
        text(s, 0.7 + i * 3.05, 2.1, 2.85, 0.6, v, size=26, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        text(s, 0.7 + i * 3.05, 2.8, 2.85, 0.4, l, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    bullets(s, 0.9, 4.0, 11.5, [
        ("Tokenise", "Lower-case the text and split into words with a regular expression."),
        ("Index and pad", "Map words to integers, keep the 10,000 most frequent, pad or trim to 200 tokens."),
    ])

    # 4 architecture
    s = new_slide(prs, "Network architecture", "Embedding, Bi-LSTM and dense classifier")
    rows = m["layers"] if m else [["embedding", "Embedding", "(None, 200, 64)", 640000],
                                   ["bilstm", "Bidirectional", "(None, 64)", 24832],
                                   ["dense", "Dense", "(None, 32)", 2080],
                                   ["dropout", "Dropout", "(None, 32)", 0],
                                   ["sentiment", "Dense", "(None, 1)", 33]]
    box(s, 0.7, 1.9, 11.9, 4.6)
    for j, hd in enumerate(["Layer", "Type", "Output shape", "Parameters"]):
        text(s, 1.1 + j * 2.9, 2.2, 2.7, 0.3, hd.upper(), size=12, bold=True, color=MUTED)
    for i, r in enumerate(rows):
        for j, v in enumerate([r[0], r[1], r[2], f"{r[3]:,}"]):
            text(s, 1.1 + j * 2.9, 2.8 + i * 0.62, 2.7, 0.4, v, size=17, color=TEXT if j else TEAL)
    text(s, 1.1, 6.0, 11, 0.3, f"Total parameters: {params}", size=14, color=MUTED)

    # 5 results
    s = new_slide(prs, "Training and results", "Adam optimiser, binary cross-entropy, dropout 0.4, early stopping")
    for i, (v, l) in enumerate([(acc, "Test accuracy"), (epochs, "Epochs"), (params, "Parameters")]):
        box(s, 0.7, 1.9 + i * 1.6, 3.2, 1.4)
        text(s, 0.7, 2.05 + i * 1.6, 3.2, 0.6, v, size=28, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        text(s, 0.7, 2.7 + i * 1.6, 3.2, 0.4, l, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    if m:
        cd = CategoryChartData()
        cd.categories = [str(i + 1) for i in range(len(m["history"]["loss"]))]
        for key, name in [("accuracy", "Train accuracy"), ("val_accuracy", "Validation accuracy"),
                          ("loss", "Train loss"), ("val_loss", "Validation loss")]:
            cd.add_series(name, m["history"][key])
        frame = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(4.3), Inches(1.8), Inches(8.3), Inches(4.9), cd)
        ch = frame.chart
        ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
        ch.legend.font.size, ch.legend.font.color.rgb = Pt(12), rgb(MUTED)
        ch.font.size, ch.font.color.rgb = Pt(12), rgb(MUTED)
        for ser, col in zip(ch.plots[0].series, [TEAL, "5EEAD4", CORAL, AMBER]):
            ser.format.line.color.rgb, ser.format.line.width, ser.smooth = rgb(col), Pt(2.5), False
        ch.value_axis.major_gridlines.format.line.color.rgb = rgb(LINE)
        ch.value_axis.format.line.fill.background()
        ch.category_axis.format.line.color.rgb = rgb(LINE)

    # 6 explainability
    s = new_slide(prs, "Making the network explain itself", "Occlusion analysis, word by word")
    for i, (n, t, d) in enumerate([("1", "Predict", "Run the full sentence through the network."),
                                    ("2", "Remove", "Delete one word at a time and predict again."),
                                    ("3", "Compare", "The change in output is that word's influence."),
                                    ("4", "Highlight", "Teal pushes positive, coral pushes negative.")]):
        x = 0.7 + i * 3.05
        box(s, x, 1.9, 2.85, 2.6)
        text(s, x + 0.3, 2.15, 2.3, 0.4, n, size=14, bold=True, color=TEAL)
        text(s, x + 0.3, 2.6, 2.3, 0.5, t, size=22, bold=True)
        text(s, x + 0.3, 3.2, 2.3, 1.2, d, size=14, color=MUTED)
    text(s, 0.9, 5.1, 11.5, 1.2, "Works with any model, needs no extra training, and the explanation reflects what the "
         "network actually computes rather than an approximation.", size=18, color=MUTED)

    # 7-10 screenshots
    for fname, title, cap in SHOTS:
        s = new_slide(prs, title, cap)
        picture_or_placeholder(s, ROOT / "screenshots" / fname, 0.7, 1.75, 11.9, 5.2)

    # 11 deployment
    s = new_slide(prs, "Deployment and tech stack", "From notebook idea to a public link")
    bullets(s, 0.9, 1.9, 5.6, [
        ("Model", "TensorFlow / Keras, saved as a single .keras file."),
        ("Interface", "Streamlit with custom styling, Plotly charts, live keystroke input."),
        ("Hosting", "Streamlit Community Cloud, deployed straight from GitHub."),
    ])
    box(s, 7.2, 1.9, 5.4, 4.6)
    text(s, 7.6, 2.2, 4.6, 0.4, "LINKS", size=12, bold=True, color=MUTED)
    text(s, 7.6, 2.8, 4.6, 0.4, "GitHub", size=14, color=MUTED)
    text(s, 7.6, 3.15, 4.6, 0.5, a.github, size=18, bold=True, color=TEAL)
    text(s, 7.6, 4.2, 4.6, 0.4, "Live demo", size=14, color=MUTED)
    text(s, 7.6, 4.55, 4.6, 0.5, a.demo, size=18, bold=True, color=TEAL)

    # 12 conclusion
    s = new_slide(prs, "Conclusion and future scope")
    bullets(s, 0.9, 1.6, 5.6, [
        ("What we built", f"A deployed sentiment network at {acc} test accuracy with live, word-level explanations."),
        ("What we learned", "Sequence models, embeddings, regularisation and how to open the black box."),
    ])
    box(s, 7.2, 1.6, 5.4, 4.2)
    text(s, 7.6, 1.9, 4.6, 0.4, "NEXT STEPS", size=12, bold=True, color=MUTED)
    text(s, 7.6, 2.4, 4.6, 3.2, ["Transformer-based model for higher accuracy", "Multi-language support",
                                 "Batch analysis of CSV files", "Fine-tuning on product and tweet data"], size=17, spacing=14)

    prs.save(ROOT / a.out)
    print("Saved", a.out)


if __name__ == "__main__":
    main()
